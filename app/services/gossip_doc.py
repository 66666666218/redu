"""把「吃瓜资料」的 markdown 渲染成 **PDF / Word**(2026-10-25,拉新物料原型)。

## 为什么单独一个模块

生成(`llm_client.compose_gossip_doc`)与**排版**是两件事:同一份 markdown 以后可能还要
渲染成公众号正文、图片卡片。把渲染放在这里,生成那边就只管"内容对不对"。

## 两条纪律

1. **必须带免责声明** —— 渲染前检查;缺了就不出文件。吃瓜物料一旦少了这句话,
   等于我们替未经证实的内容背书。
2. **不静默降级**:字体缺了 / 目录不可写就抛,别生成一份乱码或空壳 PDF 还返回成功。
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from app.utils import get_logger

logger = get_logger(__name__)

_FONT_CANDIDATES: tuple[tuple[str, str, int | None], ...] = (
    ("SimHei", r"C:\Windows\Fonts\simhei.ttf", None),
    ("MSYH", r"C:\Windows\Fonts\msyh.ttc", 0),
    ("SimSun", r"C:\Windows\Fonts\simsun.ttc", 0),
)

_DISCLAIMER_HINT = "免责声明"


class GossipDocError(RuntimeError):
    """渲染失败(字体缺失 / 目录不可写 / 内容不合规)。"""


def parse_markdown(md: str) -> tuple[str, list[tuple[str, list[str]]]]:
    """markdown → `(标题, [(小节标题, 行文本)...])`。

    只认我们自己的那套结构(`# ` 标题 + `## ` 小节),**不引第三方 markdown 库** ——
    少一个依赖,也少一处"渲染出来跟预期不一样"的意外。
    """
    title = "未命名资料"
    sections: list[tuple[str, list[str]]] = []
    cur: tuple[str, list[str]] | None = None
    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if line.startswith("# "):
            title = line[2:].strip() or title
            continue
        if line.startswith("## "):
            cur = (line[3:].strip(), [])
            sections.append(cur)
            continue
        if not line.strip():
            continue
        if cur is None:
            cur = ("正文", [])
            sections.append(cur)
        cur[1].append(line.strip())
    return title, sections


def _ensure_font() -> str:
    """注册一个中文字体,返回字体名。**找不到就抛** —— 宁可不出文件,也不出乱码。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, path, idx in _FONT_CANDIDATES:
        if not Path(path).exists():
            continue
        try:
            if idx is None:
                pdfmetrics.registerFont(TTFont(name, path))
            else:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=idx))
            return name
        except Exception:  # noqa: BLE001 - 换下一个候选
            logger.debug("字体 %s 注册失败,换下一个", path, exc_info=True)
    raise GossipDocError("找不到可用的中文字体(simhei/msyh/simsun 都不在)")


def _targets(out_dir: Path, stem: str) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    base = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", stem)[:40].strip("_") or "gossip"
    return out_dir / f"{base}_{ts}.pdf", out_dir / f"{base}_{ts}.docx"


def _check(md: str) -> None:
    if _DISCLAIMER_HINT not in (md or ""):
        raise GossipDocError(
            "内容里没有「免责声明」小节 ⇒ 拒绝出文件(吃瓜物料少了它等于替未经证实的内容背书)")


def render_pdf(md: str, path: Path) -> Path:
    """markdown → PDF(A4,中文字体)。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    font = _ensure_font()
    title, sections = parse_markdown(md)
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    h1 = ParagraphStyle("h1", fontName=font, fontSize=18, leading=26, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName=font, fontSize=13, leading=20,
                        spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=17)
    story = [Paragraph(escape(title), h1), Spacer(1, 6)]
    for head, lines in sections:
        story.append(Paragraph(escape(head), h2))
        for ln in lines:
            text = ln[2:] if ln.startswith("- ") else ln
            story.append(Paragraph("· " + escape(text), body))
    doc.build(story)
    return path


def render_docx(md: str, path: Path) -> Path:
    """markdown → Word(.docx)。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    title, sections = parse_markdown(md)
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    #: ⚠️ 中文字体要单独设 `eastAsia`,只设 `font.name` 的话中文会退回宋体/乱码
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    d.add_heading(title, level=0)
    for head, lines in sections:
        d.add_heading(head, level=1)
        for ln in lines:
            d.add_paragraph(ln[2:] if ln.startswith("- ") else ln, style="List Bullet"
                            if ln.startswith("- ") else None)
    d.save(str(path))
    return path


def build(md: str, out_dir: str | Path, stem: str = "") -> dict:
    """一份 markdown → 同名的 PDF + DOCX。返回 `{"pdf","docx","title"}`。"""
    _check(md)
    title, _ = parse_markdown(md)
    pdf_path, docx_path = _targets(Path(out_dir), stem or title)
    render_pdf(md, pdf_path)
    render_docx(md, docx_path)
    logger.info("吃瓜资料已生成:%s / %s", pdf_path.name, docx_path.name)
    return {"pdf": str(pdf_path), "docx": str(docx_path), "title": title}
